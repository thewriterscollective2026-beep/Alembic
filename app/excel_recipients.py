"""Parse an uploaded Excel (.xlsx) file and pull out all email addresses.

Standard-tolerant: scans every cell on every sheet and extracts anything that
looks like an email, so it works whether the emails are in column A, under an
'Email' header, or mixed with other columns."""
import io
import re

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")


def extract_emails_from_xlsx(data: bytes) -> list[str]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    seen, out = set(), []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell is None:
                    continue
                for match in EMAIL_RE.findall(str(cell)):
                    e = match.strip().lower()
                    if e not in seen:
                        seen.add(e)
                        out.append(e)
    wb.close()
    return out
